"""
Data Ingestion and Dynamic PostgreSQL Table Creation Service
Orchestrates file parsing, data profiling, data cleaning, schema generation,
and transactional bulk import into PostgreSQL.
"""

import os
import json
import io
from typing import Optional, Dict, Any, List
import pandas as pd
from psycopg import sql
from fastapi import HTTPException, status

import database
from backend.config import Settings
from backend.services.storage_service import StorageService
from backend.services.dataset_service import DatasetService
from backend.services.cleaning_service import CleaningService
from backend.services.schema_service import SchemaService, generate_safe_table_name
from backend.schemas.dataset import (
    DatasetProcessResponse,
    DatasetPreview,
    DatasetImportResponse
)

PREVIEW_ROWS_COUNT = 10


def load_dataset_file_to_df(stored_path: str, file_format: str) -> pd.DataFrame:
    """
    Parses an uploaded file into a Pandas DataFrame.
    Supports CSV, XLSX, JSON, and Parquet.
    """
    if not os.path.exists(stored_path):
        raise FileNotFoundError(f"Stored dataset file not found at '{stored_path}'.")

    fmt = file_format.lower()

    if fmt == "csv":
        try:
            return pd.read_csv(stored_path, encoding="utf-8")
        except UnicodeDecodeError:
            return pd.read_csv(stored_path, encoding="latin-1")
        except Exception as e:
            raise ValueError(f"Failed to parse CSV file: {str(e)}")

    elif fmt in ("xlsx", "excel"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(stored_path, read_only=True)
            sheet_names = wb.sheetnames
            if not sheet_names:
                raise ValueError("Excel file contains no worksheets.")
            # Select the first sheet
            first_sheet = sheet_names[0]
            return pd.read_excel(stored_path, sheet_name=first_sheet)
        except Exception as e:
            raise ValueError(f"Failed to parse Excel file: {str(e)}")

    elif fmt == "json":
        try:
            with open(stored_path, "r", encoding="utf-8") as f:
                raw_json = json.load(f)

            if isinstance(raw_json, list):
                # Standard array of records
                return pd.DataFrame(raw_json)
            elif isinstance(raw_json, dict):
                # Check if dictionary contains a list of records under a key (e.g. 'data' or 'records')
                for key, val in raw_json.items():
                    if isinstance(val, list) and len(val) > 0 and isinstance(val[0], dict):
                        return pd.DataFrame(val)
                # Single record dict or flat dictionary
                return pd.DataFrame([raw_json])
            else:
                raise ValueError("Unsupported JSON layout: Expected a list of objects or a tabular JSON dictionary.")
        except Exception as e:
            raise ValueError(f"Failed to parse JSON file: {str(e)}")

    elif fmt == "parquet":
        try:
            return pd.read_parquet(stored_path)
        except Exception as e:
            raise ValueError(f"Failed to parse Parquet file: {str(e)}")

    else:
        raise ValueError(f"Unsupported file format '{fmt}'.")


class IngestionService:
    """Orchestrates dataset processing and PostgreSQL table ingestion."""

    def __init__(self, settings: Settings):
        self.settings = settings
        self.storage = StorageService(settings)
        self.dataset_svc = DatasetService(settings)
        self.cleaning_svc = CleaningService()
        self.schema_svc = SchemaService()

    def _get_verified_dataset_meta(self, dataset_id: str):
        """Fetches and verifies dataset metadata and file existence."""
        dataset = self.dataset_svc.get_dataset_by_id(dataset_id)
        if not dataset:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Dataset with ID '{dataset_id}' not found."
            )

        stored_path = os.path.abspath(dataset.stored_path)
        upload_dir = os.path.abspath(self.settings.upload_dir)

        if not stored_path.startswith(upload_dir):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Security Error: Dataset storage path is outside configured upload directory."
            )

        if not os.path.exists(stored_path):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Underlying dataset file is missing from storage."
            )

        return dataset, stored_path

    def process_dataset(self, dataset_id: str) -> DatasetProcessResponse:
        """
        Parses, cleans, and profiles a dataset for admin inspection.
        Does NOT create database tables.
        """
        dataset, stored_path = self._get_verified_dataset_meta(dataset_id)

        try:
            # 1. Parse File
            raw_df = load_dataset_file_to_df(stored_path, dataset.file_format)
            if len(raw_df) == 0:
                raise ValueError("Dataset contains zero data rows.")

            # 2. Clean Data
            cleaned_df, column_map, cleaning_report = self.cleaning_svc.clean_dataframe(raw_df)

            # 3. Profile Schema & Detect Types
            schema_profiles = self.schema_svc.profile_dataset(cleaned_df, column_map)

            # 4. Generate Safe Suggested Table Name
            suggested_table = generate_safe_table_name(dataset.dataset_name)

            # 5. Format Preview (First N rows)
            preview_df = cleaned_df.head(PREVIEW_ROWS_COUNT)
            # Convert NaN/NaT to None for clean JSON serialization
            preview_records = preview_df.where(pd.notnull(preview_df), None).to_dict(orient="records")

            preview = DatasetPreview(
                total_rows=len(cleaned_df),
                total_columns=len(cleaned_df.columns),
                preview_rows=len(preview_records),
                columns=list(cleaned_df.columns),
                records=preview_records
            )

            return DatasetProcessResponse(
                status="ready_for_import",
                dataset_id=dataset_id,
                dataset_name=dataset.dataset_name,
                suggested_table_name=suggested_table,
                preview=preview,
                schema_detected=schema_profiles,
                cleaning_report=cleaning_report
            )

        except Exception as e:
            self.dataset_svc.update_dataset_status(
                dataset_id=dataset_id,
                processing_status="FAILED",
                error_message=str(e)
            )
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Dataset processing error: {str(e)}"
            )

    def import_dataset_to_database(
        self,
        dataset_id: str,
        custom_table_name: Optional[str] = None
    ) -> DatasetImportResponse:
        """
        Executes atomic, transactional PostgreSQL table creation and bulk data insertion.
        Uses the write-capable admin connection.
        """
        dataset, stored_path = self._get_verified_dataset_meta(dataset_id)

        # Update metadata status to PROCESSING
        self.dataset_svc.update_dataset_status(
            dataset_id=dataset_id,
            processing_status="PROCESSING"
        )

        try:
            # 1. Parse and Clean Data
            raw_df = load_dataset_file_to_df(stored_path, dataset.file_format)
            cleaned_df, column_map, _ = self.cleaning_svc.clean_dataframe(raw_df)

            if len(cleaned_df) == 0:
                raise ValueError("Dataset is empty after cleaning; cannot import 0 rows.")

            # 2. Determine and Validate Destination Table Name
            raw_target_name = custom_table_name.strip() if custom_table_name and custom_table_name.strip() else dataset.dataset_name
            target_table_name = generate_safe_table_name(raw_target_name)

            # 3. Detect Column Types
            schema_profiles = self.schema_svc.profile_dataset(cleaned_df, column_map)
            col_type_map = {p.normalized_name: p.detected_type for p in schema_profiles}

            # 4. Transactional PostgreSQL Table Creation & Data Insertion
            with database.get_admin_db_connection() as conn:
                with conn.cursor() as cursor:
                    # Construct CREATE TABLE using safe psycopg SQL identifiers
                    column_defs = []
                    for col_name in cleaned_df.columns:
                        col_type = col_type_map.get(col_name, "TEXT")
                        column_defs.append(
                            sql.SQL("{} {}").format(
                                sql.Identifier(col_name),
                                sql.SQL(col_type)
                            )
                        )

                    create_table_query = sql.SQL("CREATE TABLE {} ({});").format(
                        sql.Identifier(target_table_name),
                        sql.SQL(", ").join(column_defs)
                    )

                    # Execute CREATE TABLE
                    cursor.execute(create_table_query)

                    # Bulk Insert Cleaned Data using parameterized INSERT
                    cols_identifiers = sql.SQL(", ").join([sql.Identifier(c) for c in cleaned_df.columns])
                    placeholders = sql.SQL(", ").join([sql.Placeholder() for _ in cleaned_df.columns])
                    insert_query = sql.SQL("INSERT INTO {} ({}) VALUES ({});").format(
                        sql.Identifier(target_table_name),
                        cols_identifiers,
                        placeholders
                    )

                    # Prepare rows as Python native types with None for missing values
                    # Handle NaT / NaN cleanly
                    cleaned_df_sanitized = cleaned_df.where(pd.notnull(cleaned_df), None)
                    rows_to_insert = [tuple(row) for row in cleaned_df_sanitized.itertuples(index=False, name=None)]

                    # Execute batch insertion (in chunks of 1000 for high efficiency)
                    chunk_size = 1000
                    for i in range(0, len(rows_to_insert), chunk_size):
                        chunk = rows_to_insert[i:i + chunk_size]
                        cursor.executemany(insert_query, chunk)

                    # Verify row count
                    cursor.execute(sql.SQL("SELECT COUNT(*) as cnt FROM {};").format(sql.Identifier(target_table_name)))
                    imported_count = cursor.fetchone()["cnt"]
                    if imported_count != len(cleaned_df):
                        raise RuntimeError(f"Row count mismatch: Expected {len(cleaned_df)}, inserted {imported_count}.")

                # Commit transaction atomically
                conn.commit()

            # 5. Update Metadata Status to READY
            self.dataset_svc.update_dataset_status(
                dataset_id=dataset_id,
                processing_status="READY",
                table_name=target_table_name,
                row_count=len(cleaned_df),
                column_count=len(cleaned_df.columns),
                error_message=None
            )

            return DatasetImportResponse(
                status="success",
                message=f"Dataset successfully imported into PostgreSQL table '{target_table_name}' ({len(cleaned_df)} rows).",
                dataset_id=dataset_id,
                table_name=target_table_name,
                rows_imported=len(cleaned_df),
                columns_imported=len(cleaned_df.columns)
            )

        except Exception as e:
            # Clean up on failure: update metadata to FAILED
            error_details = str(e)
            self.dataset_svc.update_dataset_status(
                dataset_id=dataset_id,
                processing_status="FAILED",
                error_message=error_details
            )

            # Cleanup partial table if created
            try:
                with database.get_admin_db_connection() as cleanup_conn:
                    with cleanup_conn.cursor() as cleanup_cur:
                        cleanup_cur.execute(sql.SQL("DROP TABLE IF EXISTS {} CASCADE;").format(sql.Identifier(target_table_name)))
                    cleanup_conn.commit()
            except Exception:
                pass

            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Import failed and transaction was rolled back: {error_details}"
            )
